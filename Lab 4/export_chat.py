import json
import re
import os
import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

TRANSCRIPT_PATH = r"C:\Users\vegit\.gemini\antigravity\brain\4ff3743a-7a5e-4c7e-be57-12906d399426\.system_generated\logs\transcript_full.jsonl"
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_PATH = os.path.join(OUTPUT_DIR, "Chat_History_Lab4.docx")
PDF_PATH = os.path.join(OUTPUT_DIR, "Chat_History_Lab4.pdf")

def parse_transcript():
    turns = []
    current_user = None

    with open(TRANSCRIPT_PATH, "r", encoding="utf-8") as f:
        for line in f:
            obj = json.loads(line)
            t = obj.get("type")
            content = obj.get("content", "")
            if not content:
                continue
            if t == "USER_INPUT":
                match = re.search(r"<USER_REQUEST>(.*?)</USER_REQUEST>", content, re.DOTALL)
                req_text = match.group(1).strip() if match else content.strip()
                current_user = req_text
            elif t == "PLANNER_RESPONSE" and current_user:
                if "I am preparing the export tools" in content:
                    continue
                turns.append({
                    "prompt": current_user,
                    "response": content.strip()
                })
                current_user = None

    # Add the current export request
    if current_user and current_user == "also export the chat as a pdf or a doc":
        turns.append({
            "prompt": current_user,
            "response": "Exporting the complete chat history as both Chat_History_Lab4.pdf and Chat_History_Lab4.docx."
        })

    return turns

def generate_docx(turns):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    title = doc.add_paragraph()
    run_title = title.add_run("Lab 4: Vibe Coding — Maze Runner")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(20, 50, 110)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    run_sub = sub.add_run("LLM Interaction History & Code Review Transcript")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(90, 90, 100)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata Box
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(6)
    meta_p.paragraph_format.space_after = Pt(16)
    r1 = meta_p.add_run("Conversation ID: ")
    r1.font.bold = True
    meta_p.add_run("4ff3743a-7a5e-4c7e-be57-12906d399426\n")
    r2 = meta_p.add_run("Model: ")
    r2.font.bold = True
    meta_p.add_run("Gemini 3.8 Flash (Vibe Coding Assistant)\n")
    r3 = meta_p.add_run("Deliverable: ")
    r3.font.bold = True
    meta_p.add_run("Lab-4 Chat History Export (PDF & DOCX)\n")
    r4 = meta_p.add_run("Repo: ")
    r4.font.bold = True
    meta_p.add_run("12_maze_runner-main (Pygame)")

    doc.add_heading("Conversation Transcript", level=1)

    for i, turn in enumerate(turns, 1):
        # User Prompt Header
        p_user_head = doc.add_paragraph()
        p_user_head.paragraph_format.space_before = Pt(14)
        p_user_head.paragraph_format.space_after = Pt(4)
        r_uh = p_user_head.add_run(f"Turn {i} — User Prompt")
        r_uh.font.size = Pt(12)
        r_uh.font.bold = True
        r_uh.font.color.rgb = RGBColor(20, 80, 160)

        # User Prompt Content
        p_user = doc.add_paragraph()
        p_user.paragraph_format.left_indent = Inches(0.2)
        p_user.paragraph_format.space_after = Pt(8)
        r_uc = p_user.add_run(turn["prompt"])
        r_uc.font.italic = True
        r_uc.font.size = Pt(10.5)

        # Assistant Header
        p_asst_head = doc.add_paragraph()
        p_asst_head.paragraph_format.space_before = Pt(6)
        p_asst_head.paragraph_format.space_after = Pt(4)
        r_ah = p_asst_head.add_run("Assistant Response")
        r_ah.font.size = Pt(11)
        r_ah.font.bold = True
        r_ah.font.color.rgb = RGBColor(30, 130, 60)

        # Clean markdown links from response
        clean_resp = re.sub(r"\[(.*?)\]\(file:///[^\)]+\)", r"\1", turn["response"])
        clean_resp = re.sub(r"\[(.*?)\]\(conversation://[^\)]+\)", r"\1", clean_resp)

        for line in clean_resp.split("\n"):
            line = line.strip()
            if not line:
                continue
            p_resp = doc.add_paragraph()
            p_resp.paragraph_format.left_indent = Inches(0.2)
            p_resp.paragraph_format.space_after = Pt(2)
            if line.startswith("### "):
                r = p_resp.add_run(line[4:])
                r.font.bold = True
                r.font.size = Pt(11)
            elif line.startswith("- "):
                r = p_resp.add_run("• " + line[2:])
                r.font.size = Pt(10)
            elif re.match(r"^\d+\.\s+", line):
                r = p_resp.add_run(line)
                r.font.size = Pt(10)
            else:
                r = p_resp.add_run(line)
                r.font.size = Pt(10)

    doc.save(DOCX_PATH)
    print(f"Generated DOCX: {DOCX_PATH}")

def generate_pdf(turns):
    doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#14326e"),
        alignment=1, # Center
        spaceAfter=4
    )

    sub_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#5a5a64"),
        alignment=1,
        spaceAfter=12
    )

    meta_style = ParagraphStyle(
        "MetaText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#222230"),
        spaceAfter=10
    )

    user_head_style = ParagraphStyle(
        "UserHeader",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#1450a0"),
        spaceBefore=10,
        spaceAfter=3
    )

    user_text_style = ParagraphStyle(
        "UserText",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9.5,
        leading=13,
        leftIndent=10,
        textColor=colors.HexColor("#1a1a1a"),
        spaceAfter=6
    )

    asst_head_style = ParagraphStyle(
        "AsstHeader",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13,
        textColor=colors.HexColor("#1e823c"),
        spaceBefore=4,
        spaceAfter=2
    )

    asst_text_style = ParagraphStyle(
        "AsstText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        leftIndent=10,
        textColor=colors.HexColor("#222222"),
        spaceAfter=3
    )

    story = []

    story.append(Paragraph("Lab 4: Vibe Coding — Maze Runner", title_style))
    story.append(Paragraph("LLM Interaction History & Code Review Transcript", sub_style))

    meta_html = (
        "<b>Conversation ID:</b> 4ff3743a-7a5e-4c7e-be57-12906d399426<br/>"
        "<b>Model:</b> Gemini 3.8 Flash (Vibe Coding Assistant)<br/>"
        "<b>Deliverable:</b> Lab-4 Chat History Export (PDF & DOCX)<br/>"
        "<b>Repository:</b> 12_maze_runner-main (Pygame)"
    )
    story.append(Paragraph(meta_html, meta_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#d0d0d8"), spaceBefore=4, spaceAfter=10))

    for i, turn in enumerate(turns, 1):
        turn_story = []
        turn_story.append(Paragraph(f"Turn {i} — User Prompt", user_head_style))
        clean_prompt = html.escape(turn["prompt"]).replace("\n", "<br/>")
        turn_story.append(Paragraph(clean_prompt, user_text_style))

        turn_story.append(Paragraph("Assistant Response", asst_head_style))

        # Clean markdown formatting for ReportLab
        clean_resp = re.sub(r"\[(.*?)\]\(file:///[^\)]+\)", r"\1", turn["response"])
        clean_resp = re.sub(r"\[(.*?)\]\(conversation://[^\)]+\)", r"\1", clean_resp)

        for line in clean_resp.split("\n"):
            line = line.strip()
            if not line:
                continue
            esc_line = html.escape(line)
            # Support bold markdown **text**
            esc_line = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", esc_line)
            esc_line = re.sub(r"`(.*?)`", r"<font name='Courier'>\1</font>", esc_line)

            if esc_line.startswith("### "):
                turn_story.append(Paragraph(f"<b>{esc_line[4:]}</b>", asst_text_style))
            elif esc_line.startswith("- "):
                turn_story.append(Paragraph(f"• {esc_line[2:]}", asst_text_style))
            else:
                turn_story.append(Paragraph(esc_line, asst_text_style))

        turn_story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e4e4ec"), spaceBefore=8, spaceAfter=8))
        story.extend(turn_story)

    doc.build(story)
    print(f"Generated PDF: {PDF_PATH}")

if __name__ == "__main__":
    turns = parse_transcript()
    generate_docx(turns)
    generate_pdf(turns)
